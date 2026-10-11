"""DCF model workbook (FD-13): structure, live formulas, and a LibreOffice tie-out to engine/dcf.py."""

import io
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from engine import dcf as dcf_mod
from engine import export
from engine.dcf_workbook import dcf_workbook
from tests.test_app import fake_report


def build(mid_year=False, steady=True, peers=("P1", "P2", "P3")):
    rep = fake_report("T", list(peers))
    m, b = rep.target.metrics, rep.dcf_base
    a = dcf_mod.default_assumptions(b, m)
    res = dcf_mod.run(a, m, b, mid_year, steady_state=steady)
    d = export.DcfExport(res, a, dcf_mod.LABELS, dcf_mod.sensitivity(a, m, b, mid_year, steady_state=steady),
                         mid_year, steady)
    return rep, res, dcf_workbook(rep, d)


def rows(ws, col="B"):
    return {ws[f"{col}{r}"].value: r for r in range(1, ws.max_row + 1) if ws[f"{col}{r}"].value}


def test_sheets_and_formulas():
    _, _, data = build()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Cover", "Inputs", "DCF", "WACC", "Sensitivity", "Sources"]
    d = wb["DCF"]
    r = rows(d)
    # Projections and valuation are formulas, not pasted numbers.
    assert str(d[f"H{r['Revenue']}"].value).startswith("=")
    assert str(d[f"C{r['Implied share price']}"].value).startswith("=")
    assert "SUMPRODUCT" in wb["Sensitivity"]["D7"].value
    # Inputs are blue.
    i = rows(wb["Inputs"])
    assert wb["Inputs"][f"C{i['Terminal growth rate (perpetuity method)']}"].font.color.rgb.endswith("0000FF")
    assert wb["Sources"].max_row > 10


def test_needs_a_working_dcf():
    rep, res, _ = build()
    res.reason = "WACC n/a"
    with pytest.raises(ValueError):
        dcf_workbook(rep, export.DcfExport(res, {}, dcf_mod.LABELS, ([], [], []), False, True))


def _recalc(data: bytes, tmp_path) -> "Workbook":
    src = tmp_path / "m.xlsx"
    src.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=120)
    return load_workbook(out / "m.xlsx", data_only=True)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("mid_year,steady", [(False, True), (True, True), (False, False)])
def test_ties_to_engine(tmp_path, mid_year, steady):
    rep, res, data = build(mid_year, steady)
    wb = _recalc(data, tmp_path)
    d = wb["DCF"]
    r = rows(d)
    assert d[f"C{r['Implied share price']}"].value == pytest.approx(res.share_price.value, rel=1e-9)
    assert d[f"C{r['Enterprise value']}"].value * 1e6 == pytest.approx(res.enterprise_value.value, rel=1e-9)
    if res.exit_price.ok:
        assert d[f"D{r['Implied share price']}"].value == pytest.approx(res.exit_price.value, rel=1e-9)
    assert d[f"C{r['Status']}"].value == "Ties to the app"
    m, b = rep.target.metrics, rep.dcf_base
    a = dcf_mod.default_assumptions(b, m)
    _, _, grid = dcf_mod.sensitivity(a, m, b, mid_year, steady_state=steady)
    s = wb["Sensitivity"]
    for i, row in enumerate(grid):
        for j, f in enumerate(row):
            got = s.cell(row=7 + i, column=4 + j).value
            assert got == ("NM" if not f.ok else pytest.approx(f.value, rel=1e-9))
    for ws in wb.worksheets:
        for row_ in ws.iter_rows():
            for c in row_:
                assert not (isinstance(c.value, str) and c.value.startswith("#")), (ws.title, c.coordinate, c.value)
