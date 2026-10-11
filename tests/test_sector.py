"""Sector screen (FD-12), talking points and exports, on hand-built companies.

Five companies, EV/EBITDA 8, 10, 12, 14 and one NM:
  median = 11x, 25th = 9.5x, 75th = 12.5x (Excel QUARTILE.INC), all five in the stats except the NM.
  Company A at 8x is 8 / 11 - 1 = -27.3% vs the median.
Price 100, 52-week high 130: 100 / 130 - 1 = -23.1%.
"""

import io

import pytest

from engine import comps, export, pitch, sector
from engine.comps import CompanyRow
from tests.test_app import fake_report
from tests.test_comps import company, fig


def screen():
    rows = [company("A", 8), company("B", 14), company("C", 10), company("D", 12), company("E", None),
            company("A", 8), CompanyRow("F", "F", None, error="files 20-F/40-F (foreign filer)")]
    for r, g in zip(rows, (0.05, 0.25, 0.10, 0.20, 0.15)):
        r.metrics.figures.update({"revenue_growth_ltm": fig("growth", g, "ratio"),
                                  "equity_value": fig("equity value", 1000.0), "enterprise_value": fig("EV", 1025.0)})
    return sector.build("Widgets", rows)


def test_stats_over_every_usable_company():
    s = screen()
    st = s.stats["ev_ebitda"]
    assert st.n == 4 and (st.p25, st.median, st.p75) == (pytest.approx(9.5), 11, pytest.approx(12.5))
    assert dict(st.excluded)["E"].startswith("NM")
    assert any("A duplicates" in n for n in s.notes) and any("F: left out" in n for n in s.notes)
    assert [r.ticker for r in s.loaded] == ["A", "B", "C", "D", "E"]


def test_relative_value_and_52_week_distance():
    s = screen()
    assert s.relative["A"]["ev_ebitda"].value == pytest.approx(8 / 11 - 1)
    assert s.relative["E"]["ev_ebitda"].value is None
    assert s.loaded[0].get("price_vs_52w_high").value == pytest.approx(100 / 130 - 1)


def test_rank_skips_nm():
    assert [t for t, _ in screen().rank("ev_ebitda")] == ["B", "D", "C", "A"]


def test_fit_is_ordinary_least_squares():
    rows = [company(t, 10, ev_revenue=y) for t, y in (("A", 2.0), ("B", 4.0), ("C", 6.0))]
    for r, g in zip(rows, (0.1, 0.2, 0.3)):
        r.metrics.figures["revenue_growth_ltm"] = fig("growth", g, "ratio")
    f = sector.fit(rows, "revenue_growth_ltm", "ev_revenue")
    assert f.slope == pytest.approx(20) and f.intercept == pytest.approx(0) and f.r2 == pytest.approx(1)
    assert sector.fit(rows[:2], "revenue_growth_ltm", "ev_revenue").slope is None  # needs 3 points


def test_sector_points_cite_the_stats():
    pts = pitch.sector_points(screen())
    text = " ".join(pts)
    assert "Median LTM EV/EBITDA is 11.0x" in text and "cheapest A at 8.0x" in text and "richest B at 14.0x" in text


def test_company_points_skip_nm_and_compare_to_peers():
    rep = fake_report("T", [])
    pts = pitch.company_points(rep.target, rep.comps)
    text = " ".join(pts)
    assert "Trades at 20.0x LTM EV/EBITDA, a 82% premium to the peer median of 11.0x" in text
    assert "per share" in text  # implied range from peer EV/EBITDA


def test_sector_exports_open():
    from openpyxl import load_workbook
    from pptx import Presentation

    s = screen()
    wb = load_workbook(io.BytesIO(export.sector_workbook(s, ["a point"])))
    ws = wb["Sector screen"]
    assert ws["B5"].value == "A" and ws["J5"].value == 8  # EV/EBITDA column, a real number
    assert ws["J9"].value == "NM"
    formulas = [c.value for c in ws["J"] if isinstance(c.value, str) and c.value.startswith("=")]
    assert '=IFERROR(MEDIAN(J5:J9),"n/a")' in formulas and any("QUARTILE.INC" in f for f in formulas)
    assert "Sources" in wb.sheetnames and wb["Sources"].max_row > 5
    deck = Presentation(io.BytesIO(export.sector_deck(s, ["a point"])))
    assert len(deck.slides) >= 4


def test_report_exports_open():
    from openpyxl import load_workbook
    from pptx import Presentation

    from engine import dcf as dcf_mod

    rep = fake_report("T", [])
    m, b = rep.target.metrics, rep.dcf_base
    a = dcf_mod.default_assumptions(b, m)
    res = dcf_mod.run(a, m, b)
    d = export.DcfExport(res, a, dcf_mod.LABELS, dcf_mod.sensitivity(a, m, b, False), False, True)
    bars = rep.comps.football_field()
    wb = load_workbook(io.BytesIO(export.report_workbook(rep, d, ["x"])))
    assert {"Summary", "Financials", "EV bridge", "Comps", "Implied valuation", "DCF", "Sources"} <= set(wb.sheetnames)
    assert wb["Summary"]["B5"].value == 100.0  # share price
    deck = Presentation(io.BytesIO(export.report_deck(rep, d, bars, ["x"])))
    assert len(deck.slides) == 5


# ---- Phase 4e: ranks, residuals, filters (FD-12 additions) ------------------------------------


def test_percent_rank_matches_excel_percentrank_inc():
    assert sector.percent_rank([8, 10, 12, 14], 12) == pytest.approx(2 / 3)
    assert sector.percent_rank([8, 10, 12, 14], 8) == 0 and sector.percent_rank([8, 10, 12, 14], 14) == 1
    assert sector.percent_rank([1, 1, 2], 1) == 0  # ties share the lower rank
    assert sector.percent_rank([5], 5) is None


def test_ranks_skip_nm():
    r = sector.ranks(screen(), "ev_ebitda")
    assert r["A"].value == 0 and r["B"].value == 1 and r["C"].value == pytest.approx(1 / 3)
    assert r["E"].value is None and "EV/EBITDA NM" in r["E"].reason


def test_residuals_by_hand():
    """(10%, 2x), (20%, 4x), (30%, 7x): slope 25, intercept -0.667; fitted 1.83 / 4.33 / 6.83."""
    rows = [company(t, 10, ev_revenue=y) for t, y in (("A", 2.0), ("B", 4.0), ("C", 7.0))]
    for r, g in zip(rows, (0.1, 0.2, 0.3)):
        r.metrics.figures["revenue_growth_ltm"] = fig("growth", g, "ratio")
    f, res = sector.residuals(rows, "revenue_growth_ltm", "ev_revenue")
    assert f.slope == pytest.approx(25)
    assert res[0].ticker == "B" and {z.ticker for z in res} == {"A", "B", "C"}
    b = res[0]
    assert b.gap == pytest.approx(4 - 13 / 3) and b.gap_pct == pytest.approx(4 / (13 / 3) - 1)
    assert sector.residuals(rows[:2], "revenue_growth_ltm", "ev_revenue")[1] == []


def test_filter_and_subset_recompute_stats():
    s = screen()  # growth A 5%, B 25%, C 10%, D 20%, E 15%; EV/EBITDA 8, 14, 10, 12, NM
    kept, dropped = sector.filter_tickers(s, {"revenue_growth_ltm": (0.12, None)})
    assert kept == ["B", "D", "E"] and dropped == ["A", "C"]
    kept, dropped = sector.filter_tickers(s, {"ev_ebitda": (None, 13.0)})
    assert kept == ["A", "C", "D"] and "E" in dropped  # NM can't be shown to pass
    assert sector.filter_tickers(s, {"ev_ebitda": (None, None)})[1] == []
    sub = sector.subset(s, ["A", "C", "D"], "Widgets (filtered)")
    assert sub.name == "Widgets (filtered)" and sub.stats["ev_ebitda"].median == 10
    assert s.stats["ev_ebitda"].median == 11  # the original screen is unchanged


def test_below_the_line_point_needs_a_real_fit():
    s = screen()
    pts = " ".join(pitch.sector_points(s))
    assert "Furthest below the line" not in pts  # every EV/Revenue is 10x: no fit at all
    for r, y in zip(s.loaded, (2.0, 10.0, 3.0, 8.0, 6.0)):  # growth 5/25/10/20/15%: about 40x growth
        r.metrics.figures["ev_revenue"] = fig("EV/Revenue", y, "x")
    pts = " ".join(pitch.sector_points(s))
    assert "Furthest below the line (cheapest for its growth): C at 3.0x" in pts


def test_pitch_deck_and_workbook_sheets():
    from openpyxl import load_workbook
    from pptx import Presentation

    from engine import dcf as dcf_mod
    from engine import thesis

    rep = fake_report("T", [])
    m, b = rep.target.metrics, rep.dcf_base
    a = dcf_mod.default_assumptions(b, m)
    res = dcf_mod.run(a, m, b)
    d = export.DcfExport(res, a, dcf_mod.LABELS, dcf_mod.sensitivity(a, m, b, False), False, True)
    x = thesis.build_extras(rep, a)
    bars = rep.comps.football_field()
    s = screen()
    full = Presentation(io.BytesIO(export.pitch_deck(rep, s, d, bars, x, ["sector point"], ["company point"])))
    titles = [next(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame).split("\n")[0] for sl in full.slides]
    assert titles[0].startswith("Target Co (T)") and titles[1] == "Valuation at a glance"
    assert titles[2] == "Widgets: sector overview" and "T vs peers: scorecard" in titles
    assert titles.index("Trading comparables") < titles.index("Discounted cash flow") < \
        titles.index("Valuation summary (football field)") < titles.index("Appendix: methodology and data notes")
    assert len(full.slides) == 15
    alone = Presentation(io.BytesIO(export.pitch_deck(rep, None, d, bars, x)))
    assert len(alone.slides) == 10  # no sector overview, table, scatter or the two ranking slides
    wb = load_workbook(io.BytesIO(export.report_workbook(rep, d, ["x"], x)))
    assert {"Valuation summary", "Scorecard", "What moves the DCF", "Peer size check"} <= set(wb.sheetnames)
    assert wb.sheetnames[-1] == "Sources"
