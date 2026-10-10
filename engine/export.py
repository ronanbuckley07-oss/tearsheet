"""Excel and PowerPoint exports for a pitch team.

Excel: real numbers (USD millions, multiples, percentages) with number formats,
so they can be re-sorted and charted; summary rows are live Excel formulas
(MEDIAN, QUARTILE.INC: the same method as FD-7 / FD-10 Q11) so every statistic
can be audited in the sheet; and a Sources sheet listing the formula, XBRL
facts and flags behind each figure (CLAUDE.md rule 2).

PowerPoint: 16:9 slides with native (editable) tables and charts, so a team
can restyle them to its own template. Every slide carries a source line.

No finance logic: everything here formats figures the engine already computed.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date

from engine.comps import CONTEXT, MULTIPLES, VALUATION_BASES
from engine.periods import Figure
from engine.report import Report
from engine.sector import COLUMNS, LABELS as SECTOR_LABELS, RELATIVE, UNITS as SECTOR_UNITS, Sector
from engine.sources import fact_rows, other_sources

COMP_LABELS = {"equity_value": "Market cap ($m)", "enterprise_value": "EV ($m)", "ev_revenue": "EV/Revenue",
               "ev_ebitda": "EV/EBITDA", "pe": "P/E", "revenue_growth_ltm": "Revenue growth (LTM)",
               "ebitda_margin": "EBITDA margin", "net_debt_ebitda": "Net debt / EBITDA"}
COMP_UNITS = {"equity_value": "USD", "enterprise_value": "USD", "ev_revenue": "x", "ev_ebitda": "x", "pe": "x",
              "revenue_growth_ltm": "ratio", "ebitda_margin": "ratio", "net_debt_ebitda": "x"}
COMP_COLUMNS = ("equity_value", "enterprise_value", *MULTIPLES, *CONTEXT)


@dataclass
class DcfExport:
    """What the app's DCF section computed, on the user's current assumptions."""

    result: object  # dcf.Result
    assumptions: dict[str, Figure]
    labels: dict[str, str]
    sensitivity: tuple[list[float], list[float], list[list[Figure]]]
    mid_year: bool
    steady_state: bool


# ---- shared helpers -------------------------------------------------------------------


def cell_value(f: Figure | None):
    """Number in sheet units (USD and shares in millions), or "NM" / "n/a"."""
    if f is None:
        return "n/a"
    if f.value is None:
        return "NM" if f.nm else "n/a"
    return f.value / 1e6 if f.unit in ("USD", "shares") else f.value


NUMBER_FORMATS = {"USD": '#,##0;(#,##0)', "shares": '#,##0', "USD/share": '$#,##0.00', "x": '0.0"x"',
                  "ratio": '0.0%', "": 'General'}


def _why(f: Figure | None) -> str:
    if f is None or f.value is not None:
        return ""
    return f"{'NM' if f.nm else 'n/a'}: {f.reason or ''}"


def _fmt(f: Figure | None) -> str:
    """Display text for slides (same rules as the app)."""
    if f is None or f.value is None:
        return "NM" if f is not None and f.nm else "n/a"
    v, u = f.value, f.unit
    sign = "-" if v < 0 else ""
    if u == "USD":
        a = abs(v)
        return f"{sign}${a / 1e9:,.1f}bn" if a >= 1e9 else f"{sign}${a / 1e6:,.0f}m"
    if u == "shares":
        return f"{v / 1e6:,.0f}m"
    if u == "USD/share":
        return f"{sign}${abs(v):,.2f}"
    if u == "x":
        return f"{v:.1f}x"
    if u == "ratio":
        return f"{v:.1%}"
    return f"{v:,.2f}"


def _num(value: float | None, unit: str) -> str:
    return _fmt(None if value is None else Figure("", value, unit, ""))


# ---- Excel ----------------------------------------------------------------------------


class _Book:
    """Thin openpyxl wrapper: titled sheets, a header style, and a sources log."""

    def __init__(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.Font, self.Fill, self.Align = Font, PatternFill, Alignment
        self.head_font = Font(bold=True, color="FFFFFF")
        self.head_fill = PatternFill("solid", fgColor="1F3864")
        self.stat_font = Font(italic=True)
        self.sources: list[list] = []

    def sheet(self, title: str, heading: str, sub: str = ""):
        ws = self.wb.create_sheet(title[:31])
        ws["A1"] = heading
        ws["A1"].font = self.Font(bold=True, size=14)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = self.Font(italic=True, color="595959")
        ws.sheet_view.showGridLines = False
        return ws

    def header(self, ws, row: int, values: list[str], widths: list[int] | None = None) -> None:
        for j, v in enumerate(values, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.font, c.fill = self.head_font, self.head_fill
            c.alignment = self.Align(wrap_text=True, vertical="center")
        if widths:
            from openpyxl.utils import get_column_letter
            for j, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(j)].width = w

    def put(self, ws, row: int, col: int, f: Figure | None, unit: str | None = None, *,
            section: str = "", label: str = "", cik: int | None = None) -> None:
        c = ws.cell(row=row, column=col, value=cell_value(f))
        if f is not None and f.value is not None:
            c.number_format = NUMBER_FORMATS.get(unit or f.unit, "General")
        else:
            c.alignment = self.Align(horizontal="right")
        if f is not None and (section or label):
            self.log(section, label, f, cik)

    def log(self, section: str, label: str, f: Figure, cik: int | None) -> None:
        facts = fact_rows(f, cik)
        refs = "; ".join(f"{r['']}{r['tag']} {r['period']} {r['form']} {r['accession']} filed {r['filed']}"
                         for r in facts)
        link = next((r["link"] for r in facts if r["link"]), "")
        self.sources.append([section, label, cell_value(f), f.unit, f.formula, _why(f), refs,
                             " | ".join(other_sources(f)), " | ".join(f.all_notes()), link])

    def sources_sheet(self) -> None:
        ws = self.sheet("Sources", "Where every number comes from",
                        "Money and shares in millions. Facts: sign, XBRL tag, period, form, accession, filed date.")
        cols = ["Section", "Figure", "Value", "Unit", "Formula", "NM / n/a reason", "XBRL facts",
                "Other sources", "Notes and flags", "Filing link"]
        self.header(ws, 4, cols, [16, 34, 14, 10, 50, 40, 80, 50, 60, 40])
        for i, row in enumerate(self.sources, 5):
            for j, v in enumerate(row, 1):
                ws.cell(row=i, column=j, value=v)
            if row[9]:
                ws.cell(row=i, column=10).hyperlink = row[9]
        ws.freeze_panes = "C5"

    def save(self) -> bytes:
        self.wb.calculation.fullCalcOnLoad = True  # formulas compute when the file opens
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def _stat_rows(book: _Book, ws, first: int, last: int, start_row: int, cols: dict[int, str],
               label_col: int = 1, label_prefix: str = "") -> int:
    """Median / mean / 25th / 75th rows as Excel formulas over rows first..last (text NM/n/a ignored)."""
    from openpyxl.utils import get_column_letter

    r = start_row
    for label, fn in (("Median", "MEDIAN({})"), ("Mean", "AVERAGE({})"),
                      ("25th percentile", "QUARTILE.INC({},1)"), ("75th percentile", "QUARTILE.INC({},3)")):
        ws.cell(row=r, column=label_col, value=f"{label_prefix}{label}").font = book.stat_font
        for col, unit in cols.items():
            L = get_column_letter(col)
            c = ws.cell(row=r, column=col, value=f'=IFERROR({fn.format(f"{L}{first}:{L}{last}")},"n/a")')
            c.number_format = NUMBER_FORMATS[unit]
            c.font = book.stat_font
        r += 1
    return r


def report_workbook(rep: Report, dcf: DcfExport | None = None, points: list[str] | None = None) -> bytes:
    book = _Book()
    t = rep.target
    m = t.metrics
    F = m.figures
    src = (f"Source: SEC EDGAR XBRL filings (LTM through {m.ltm_end}; latest filing {t.latest_filing}); "
           f"share price {F['price'].formula}, Yahoo Finance. $ in millions except per share.")

    # Summary
    ws = book.sheet("Summary", f"{t.name} ({t.ticker})", src)
    book.header(ws, 4, ["Item", "Value", "Note"], [40, 18, 90])
    snap = [("Share price", F["price"]), ("Shares outstanding (m, cover page)", F["shares_outstanding"]),
            ("Equity value ($m)", F["equity_value"]), ("Enterprise value ($m)", F["enterprise_value"]),
            ("LTM revenue ($m)", F["ltm_revenue"]), ("LTM EBITDA ($m)", F["ltm_ebitda"]),
            ("EBITDA margin (LTM)", F["ebitda_margin"]), ("Revenue growth (LTM)", F["revenue_growth_ltm"]),
            ("EV/Revenue (LTM)", F["ev_revenue"]), ("EV/EBITDA (LTM)", F["ev_ebitda"]), ("P/E (LTM)", F["pe"]),
            ("Net debt / EBITDA", F["net_debt_ebitda"]), ("52-week high", t.high_52w), ("52-week low", t.low_52w)]
    r = 5
    for label, f in snap:
        ws.cell(row=r, column=1, value=label)
        book.put(ws, r, 2, f, section="Summary", label=label, cik=t.cik)
        ws.cell(row=r, column=3, value=_why(f))
        r += 1
    if dcf is not None and not dcf.result.reason:
        ws.cell(row=r, column=1, value="DCF share price (app assumptions)")
        book.put(ws, r, 2, dcf.result.share_price, section="Summary", label="DCF share price", cik=t.cik)
        r += 1
    if points:
        r += 1
        ws.cell(row=r, column=1, value="Talking points (computed, check against the tables)").font = book.Font(bold=True)
        for p in points:
            r += 1
            ws.cell(row=r, column=1, value=f"• {p}")

    # Financials (FD-4)
    ws = book.sheet("Financials", "Financial summary ($m)", src)
    periods = [(f"FY{e.year} ({e})", m.fiscal_years[e]) for e in sorted(m.fiscal_years)]
    periods.append((f"LTM ({m.ltm_end})", {
        "revenue": F["ltm_revenue"], "operating_income": F["ltm_operating_income"], "d_and_a": F["ltm_d_and_a"],
        "ebitda": F["ltm_ebitda"], "ebitda_margin": F["ebitda_margin"], "net_income": F["ltm_net_income"],
        "sbc": F["ltm_sbc"]}))
    book.header(ws, 4, ["", *[p for p, _ in periods]], [36] + [20] * len(periods))
    lines = [("Revenue", "revenue"), ("Operating income (EBIT)", "operating_income"), ("D&A", "d_and_a"),
             ("EBITDA (FD-1)", "ebitda"), ("EBITDA margin", "ebitda_margin"), ("Net income", "net_income"),
             ("Stock-based comp (not added back)", "sbc")]
    for i, (label, k) in enumerate(lines, 5):
        ws.cell(row=i, column=1, value=label)
        for j, (p, d) in enumerate(periods, 2):
            book.put(ws, i, j, d.get(k), section="Financials", label=f"{label}, {p}", cik=t.cik)

    # EV bridge (FD-2, FD-3)
    ws = book.sheet("EV bridge", "Equity value to enterprise value ($m)", src)
    book.header(ws, 4, ["", "Value", "As of", "Note"], [44, 18, 14, 80])
    bridge = [("Share price ($)", F["price"]), ("× Shares outstanding (m)", F["shares_outstanding"]),
              ("= Equity value", F["equity_value"]), ("+ Total debt incl. finance leases", F["total_debt"]),
              ("+ Preferred stock", F["preferred_stock"]), ("+ Noncontrolling interest", F["noncontrolling_interest"]),
              ("− Cash and equivalents", F["cash"]), ("− Short-term investments", F["short_term_investments"]),
              ("= Enterprise value", F["enterprise_value"]), ("Net debt", F["net_debt"])]
    for i, (label, f) in enumerate(bridge, 5):
        ws.cell(row=i, column=1, value=label)
        book.put(ws, i, 2, f, section="EV bridge", label=label, cik=t.cik)
        ws.cell(row=i, column=3, value=str(f.as_of or ""))
        ws.cell(row=i, column=4, value=_why(f))

    # Comps (FD-5..FD-7)
    c = rep.comps
    if c.peers:
        ws = book.sheet("Comps", "Trading comps (LTM)",
                        src + " Peer statistics are Excel formulas over the peer rows only (target excluded, "
                              "NM and n/a ignored), QUARTILE.INC as in FD-7.")
        cols = ["Company", "Ticker", "LTM end", *[COMP_LABELS[k] for k in COMP_COLUMNS], "Status"]
        book.header(ws, 4, cols, [30, 9, 12] + [14] * len(COMP_COLUMNS) + [60])
        ws.cell(row=5, column=1, value=f"{t.name} (target)").font = book.Font(bold=True)
        rows = [t, *c.peers]
        for i, row in enumerate(rows, 5):
            if i > 5:
                ws.cell(row=i, column=1, value=row.name)
            ws.cell(row=i, column=2, value=row.ticker)
            if row.metrics is None:
                ws.cell(row=i, column=4 + len(COMP_COLUMNS), value=row.error)
                continue
            ws.cell(row=i, column=3, value=str(row.metrics.ltm_end))
            reasons = []
            for j, k in enumerate(COMP_COLUMNS, 4):
                f = row.get(k)
                book.put(ws, i, j, f, COMP_UNITS[k], section="Comps", label=f"{row.ticker}: {COMP_LABELS[k]}",
                         cik=row.cik)
                if _why(f):
                    reasons.append(f"{COMP_LABELS[k]} {_why(f)}")
            ws.cell(row=i, column=4 + len(COMP_COLUMNS), value="; ".join(reasons))
        last = 4 + len(rows)
        _stat_rows(book, ws, 6, last, last + 2,
                   {j: COMP_UNITS[k] for j, k in enumerate(COMP_COLUMNS, 4)}, label_prefix="Peers: ")
        ws.freeze_panes = "C5"

        # Implied valuation (FD-8)
        ws = book.sheet("Implied valuation", "Implied valuation from peer multiples (FD-8)", src)
        book.header(ws, 4, ["Method", "Peer 25th pct", "Peer median", "Peer 75th pct", "Target LTM metric ($m)",
                            "Implied EV low", "Implied EV mid", "Implied EV high", "Price low", "Price mid",
                            "Price high", "Note"], [16, 13, 13, 13, 18, 15, 15, 15, 12, 12, 12, 60])
        for i, (key, rr) in enumerate(c.implied.items(), 5):
            ws.cell(row=i, column=1, value=COMP_LABELS[key])
            for j, v in enumerate((rr.low_multiple, rr.mid_multiple, rr.high_multiple), 2):
                cc = ws.cell(row=i, column=j, value=v if v is not None else "n/a")
                cc.number_format = NUMBER_FORMATS["x"]
            book.put(ws, i, 5, rr.target_metric)
            for j, f in enumerate(rr.enterprise_value, 6):
                book.put(ws, i, j, f)
            for j, (f, point) in enumerate(zip(rr.share_price, ("25th pct", "median", "75th pct")), 9):
                book.put(ws, i, j, f, section="Implied valuation", label=f"{COMP_LABELS[key]} {point} price",
                         cik=t.cik)
            ws.cell(row=i, column=12, value=rr.reason or _why(rr.share_price[0]))

    # DCF (FD-11)
    if dcf is not None and not dcf.result.reason:
        _dcf_sheet(book, dcf, t, F, src)

    book.sources_sheet()
    return book.save()


def _dcf_sheet(book: _Book, d: DcfExport, t, F, src: str) -> None:
    res = d.result
    ws = book.sheet("DCF", "Unlevered DCF on the app's current assumptions (FD-11)",
                    src + (" Mid-year discounting." if d.mid_year else " End-of-year discounting.")
                    + (" Steady-state terminal year (capex capped at D&A)." if d.steady_state else ""))
    book.header(ws, 4, ["Assumption", "Value", "Basis"], [40, 14, 90])
    r = 5
    for k, f in d.assumptions.items():
        ws.cell(row=r, column=1, value=d.labels.get(k, k))
        book.put(ws, r, 2, f, section="DCF", label=d.labels.get(k, k), cik=t.cik)
        ws.cell(row=r, column=3, value=f.formula if f.ok else _why(f))
        r += 1
    r += 1
    heads = ["($m)", *[f"Year {y.n}" for y in res.years]]
    book.header(ws, r, heads)
    lines = [("Revenue", "revenue"), ("Growth", "growth"), ("EBIT", "ebit"), ("− Taxes on EBIT", "taxes"),
             ("= NOPAT", "nopat"), ("+ D&A", "d_and_a"), ("− Capex", "capex"), ("− Increase in NWC", "nwc_change"),
             ("= Unlevered free cash flow", "ufcf"), ("Discount factor", "discount_factor"),
             ("PV of free cash flow", "pv")]
    for label, attr in lines:
        r += 1
        ws.cell(row=r, column=1, value=label)
        for j, y in enumerate(res.years, 2):
            v = getattr(y, attr)
            if attr in ("taxes", "capex", "nwc_change"):
                v = -v
            if attr == "growth":
                cc = ws.cell(row=r, column=j, value=v)
                cc.number_format = NUMBER_FORMATS["ratio"]
            elif attr == "discount_factor":
                cc = ws.cell(row=r, column=j, value=v)
                cc.number_format = "0.000"
            else:
                cc = ws.cell(row=r, column=j, value=v / 1e6)
                cc.number_format = NUMBER_FORMATS["USD"]
    r += 2
    book.header(ws, r, ["Output", "Value"])
    outs = [("PV of free cash flows ($m)", res.pv_fcf), ("PV of terminal value ($m)", res.pv_terminal),
            ("Enterprise value ($m)", res.enterprise_value), ("Equity value ($m)", res.equity_value),
            ("Share price", res.share_price), ("WACC", res.wacc), ("Cost of equity", res.cost_of_equity)]
    for label, f in outs:
        r += 1
        ws.cell(row=r, column=1, value=label)
        book.put(ws, r, 2, f, section="DCF", label=label, cik=t.cik)
    ws_, gs, grid = d.sensitivity
    if grid:
        r += 2
        ws.cell(row=r, column=1, value="Share price: WACC (rows) × terminal growth (columns)").font = book.Font(bold=True)
        r += 1
        book.header(ws, r, ["WACC \\ g", *[f"{g:.2%}" for g in gs]])
        for w, row in zip(ws_, grid):
            r += 1
            cc = ws.cell(row=r, column=1, value=w)
            cc.number_format = "0.00%"
            for j, f in enumerate(row, 2):
                book.put(ws, r, j, f)


def sector_workbook(s: Sector, points: list[str] | None = None) -> bytes:
    book = _Book()
    rows = s.loaded
    as_of = max((r.get("price").as_of for r in rows if r.get("price").as_of), default=None)
    src = (f"Source: SEC EDGAR XBRL filings (each company's latest 10-K/10-Q; LTM end per row); prices from Yahoo "
           f"Finance as of {as_of}. $ in millions. Statistics are Excel formulas over every company (FD-12).")
    ws = book.sheet("Sector screen", f"{s.name}: sector screen", src)
    cols = ["Company", "Ticker", "LTM end", *[_sector_head(k) for k in COLUMNS],
            *[f"{SECTOR_LABELS[k]} vs median" for k in RELATIVE], "NM / n/a reasons"]
    book.header(ws, 4, cols, [30, 9, 12] + [14] * (len(COLUMNS) + len(RELATIVE)) + [80])
    for i, row in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=row.name)
        ws.cell(row=i, column=2, value=row.ticker)
        ws.cell(row=i, column=3, value=str(row.metrics.ltm_end))
        reasons = []
        for j, k in enumerate(COLUMNS, 4):
            f = row.get(k)
            book.put(ws, i, j, f, SECTOR_UNITS[k], section="Sector", label=f"{row.ticker}: {SECTOR_LABELS[k]}",
                     cik=row.cik)
            if _why(f):
                reasons.append(f"{SECTOR_LABELS[k]} {_why(f)}")
        for j, k in enumerate(RELATIVE, 4 + len(COLUMNS)):
            book.put(ws, i, j, s.relative[row.ticker][k], "ratio")
        ws.cell(row=i, column=4 + len(COLUMNS) + len(RELATIVE), value="; ".join(reasons))
    last = 4 + len(rows)
    _stat_rows(book, ws, 5, last, last + 2, {j: SECTOR_UNITS[k] for j, k in enumerate(COLUMNS, 4)},
               label_prefix="Sector: ")
    ws.freeze_panes = "C5"
    excluded = [r for r in s.rows if r.metrics is None]
    r = last + 8
    for row in excluded:
        ws.cell(row=r, column=1, value=f"{row.ticker}: left out, {row.error}")
        r += 1
    if points:
        r += 1
        ws.cell(row=r, column=1, value="Talking points (computed, check against the table)").font = book.Font(bold=True)
        for p in points:
            r += 1
            ws.cell(row=r, column=1, value=f"• {p}")
    book.sources_sheet()
    return book.save()


def _sector_head(k: str) -> str:
    return SECTOR_LABELS[k] + (" ($m)" if SECTOR_UNITS[k] == "USD" else "")


# ---- PowerPoint -------------------------------------------------------------------------

NAVY, GRAY_TXT, BLUE, ORANGE, LIGHT = "1F3864", "595959", "2A78D6", "EB6834", "D9E2F3"


class _Deck:
    def __init__(self, footer: str):
        from pptx import Presentation
        from pptx.util import Inches

        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.footer = footer

    def slide(self, title: str, subtitle: str = ""):
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.7)).text_frame
        tb.text = title
        p = tb.paragraphs[0]
        p.alignment = 1  # left
        p.runs[0].font.size, p.runs[0].font.bold = Pt(26), True
        p.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
        if subtitle:
            p2 = tb.add_paragraph()
            p2.text = subtitle
            p2.alignment = 1
            p2.runs[0].font.size = Pt(13)
            p2.runs[0].font.color.rgb = RGBColor.from_string(GRAY_TXT)
        line = s.shapes.add_connector(1, Inches(0.5), Inches(1.25), Inches(12.83), Inches(1.25))
        line.line.color.rgb = RGBColor.from_string(NAVY)
        ft = s.shapes.add_textbox(Inches(0.5), Inches(6.95), Inches(12.3), Inches(0.4)).text_frame
        ft.word_wrap = True
        ft.text = self.footer
        ft.paragraphs[0].runs[0].font.size = Pt(9)
        ft.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(GRAY_TXT)
        return s

    def table(self, s, data: list[list[str]], left, top, width, height, bold_rows: set[int] = frozenset(),
              first_col_width=None, font=10):
        from pptx.dml.color import RGBColor
        from pptx.util import Pt

        nrows, ncols = len(data), len(data[0])
        shape = s.shapes.add_table(nrows, ncols, left, top, width, height)
        tbl = shape.table
        if first_col_width is not None:
            rest = (width - first_col_width) // max(ncols - 1, 1)
            tbl.columns[0].width = first_col_width
            for j in range(1, ncols):
                tbl.columns[j].width = rest
        for i, row in enumerate(data):
            for j, v in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = str(v)
                para = cell.text_frame.paragraphs[0]
                para.alignment = 1 if j == 0 else 3  # left / right
                run = para.runs[0] if para.runs else para.add_run()
                run.font.size = Pt(font)
                run.font.bold = i == 0 or i in bold_rows
                cell.margin_left = cell.margin_right = Pt(4)
                cell.margin_top = cell.margin_bottom = Pt(2)
                if i == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(NAVY)
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
                elif i in bold_rows:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(LIGHT)
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF" if i % 2 else "F2F2F2")
        return tbl

    def bullets(self, s, lines: list[str], left, top, width, height, size=14):
        from pptx.util import Pt

        tf = s.shapes.add_textbox(left, top, width, height).text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {line}"
            p.runs[0].font.size = Pt(size)
            p.space_after = Pt(8)

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def _style_chart(chart, legend: bool = True):
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt

    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
    chart.font.size = Pt(11)


def report_deck(rep: Report, dcf: DcfExport | None = None, football: list[dict] | None = None,
                points: list[str] | None = None) -> bytes:
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches, Pt

    t = rep.target
    m = t.metrics
    F = m.figures
    deck = _Deck(f"Source: SEC EDGAR XBRL filings, LTM through {m.ltm_end} (latest filing {t.latest_filing}); "
                 f"share price {F['price'].formula} (Yahoo Finance). NM = not meaningful. Built with Tearsheet.")

    # 1. Overview: key stats + talking points
    s = deck.slide(f"{t.name} ({t.ticker})", f"{t.sic_description or ''} · Company profile · {date.today():%B %d, %Y}")
    stats = [["Market data", ""], ["Share price", _fmt(F["price"])],
             ["52-week range", f"{_fmt(t.low_52w)} - {_fmt(t.high_52w)}" if t.low_52w and t.low_52w.ok else "n/a"],
             ["Equity value", _fmt(F["equity_value"])], ["Enterprise value", _fmt(F["enterprise_value"])],
             ["LTM revenue", _fmt(F["ltm_revenue"])], ["LTM EBITDA", _fmt(F["ltm_ebitda"])],
             ["EBITDA margin", _fmt(F["ebitda_margin"])], ["Revenue growth (LTM)", _fmt(F["revenue_growth_ltm"])],
             ["EV/EBITDA (LTM)", _fmt(F["ev_ebitda"])], ["P/E (LTM)", _fmt(F["pe"])],
             ["Net debt / EBITDA", _fmt(F["net_debt_ebitda"])]]
    deck.table(s, stats, Inches(0.5), Inches(1.5), Inches(4.6), Inches(5.0), first_col_width=Inches(2.6), font=12)
    if points:
        deck.bullets(s, points[:7], Inches(5.5), Inches(1.5), Inches(7.3), Inches(5.2), size=14)

    # 2. Financial summary: table + revenue/EBITDA chart
    s = deck.slide("Financial summary", "Last two fiscal years and LTM. EBITDA = operating income + D&A (FD-1).")
    periods = [(f"FY{e.year}", m.fiscal_years[e]) for e in sorted(m.fiscal_years)]
    periods.append((f"LTM {m.ltm_end:%b-%y}", {"revenue": F["ltm_revenue"], "operating_income": F["ltm_operating_income"],
                                               "ebitda": F["ltm_ebitda"], "ebitda_margin": F["ebitda_margin"],
                                               "net_income": F["ltm_net_income"]}))
    lines = [("Revenue", "revenue"), ("EBIT", "operating_income"), ("EBITDA", "ebitda"),
             ("EBITDA margin", "ebitda_margin"), ("Net income", "net_income")]
    data = [["", *[p for p, _ in periods]]] + [[label, *[_fmt(d.get(k)) for _, d in periods]] for label, k in lines]
    deck.table(s, data, Inches(0.5), Inches(1.6), Inches(5.6), Inches(2.6), first_col_width=Inches(1.8), font=12)
    cd = CategoryChartData()
    cd.categories = [p for p, _ in periods]
    for name, k in (("Revenue", "revenue"), ("EBITDA", "ebitda")):
        cd.add_series(name, [round(d[k].value / 1e9, 2) if d.get(k) is not None and d[k].ok else None
                             for _, d in periods])
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.5), Inches(1.5), Inches(6.3), Inches(5.2),
                            cd).chart
    _style_chart(ch)
    for series, color in zip(ch.series, (BLUE, ORANGE)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = '"$"#,##0.0"bn"', False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.visible = False

    # 3. Comps
    c = rep.comps
    if c.peers:
        s = deck.slide("Trading comparables", f"LTM multiples; peer statistics exclude the target and NM values (FD-7). "
                                              f"{len(c.peers)} peers.")
        head = ["Company", "Mkt cap", "EV", "EV/Rev", "EV/EBITDA", "P/E", "Rev growth", "EBITDA mgn"]
        keys = ("equity_value", "enterprise_value", "ev_revenue", "ev_ebitda", "pe", "revenue_growth_ltm",
                "ebitda_margin")
        data = [head]
        for row in c.peers:
            if row.metrics is not None:
                data.append([row.name[:28], *[_fmt(row.get(k)) for k in keys]])
        stat_start = len(data)
        for label, attr in (("Peer median", "median"), ("Peer 25th pct", "p25"), ("Peer 75th pct", "p75")):
            line = [label, "", ""]
            for k in keys[2:]:
                st = c.stats.get(k)
                v = getattr(st, attr) if st else None
                line.append(_num(v, COMP_UNITS[k]))
            data.append(line)
        data.append([f"{t.name[:24]} (target)", *[_fmt(t.get(k)) for k in keys]])
        n = len(data)
        deck.table(s, data, Inches(0.5), Inches(1.5), Inches(12.3), Inches(min(5.2, 0.32 * n)),
                   bold_rows={*range(stat_start, n)}, first_col_width=Inches(3.3), font=11 if n < 16 else 9)

    # 4. Football field
    bars = [b for b in (football or []) if b["low"] is not None and b["high"] is not None
            and b["label"] != "Current price"]
    price = F["price"].value
    if bars:
        s = deck.slide("Valuation summary (football field)",
                       f"Implied share price ranges vs current price {_fmt(F['price'])}.")
        cd = CategoryChartData()
        labels = [b["label"] for b in reversed(bars)]
        if price:
            labels.append(f"Current price ({_fmt(F['price'])})")
        cd.categories = labels
        lows = [b["low"] for b in reversed(bars)] + ([price * 0.996] if price else [])
        spans = [b["high"] - b["low"] for b in reversed(bars)] + ([price * 0.008] if price else [])
        cd.add_series("base", [round(v, 2) for v in lows])
        cd.add_series("range", [round(v, 2) for v in spans])
        ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3),
                                cd).chart
        _style_chart(ch, legend=False)
        ch.series[0].format.fill.background()
        rng = ch.series[1]
        rng.format.fill.solid()
        rng.format.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        for i, b in enumerate(reversed(bars)):
            pt = rng.points[i]
            if b["label"] == "52-week range":
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = RGBColor.from_string("8A8984")
            pt.data_label.has_text_frame = True
            pt.data_label.text_frame.text = f"${b['low']:,.0f} - ${b['high']:,.0f}"
            pt.data_label.position = XL_LABEL_POSITION.INSIDE_BASE
            pt.data_label.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
            pt.data_label.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        if price:
            pt = rng.points[len(bars)]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(ORANGE)
        ch.plots[0].gap_width = 60
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '"$"#,##0', False
        lo_all = min(lows)
        ch.value_axis.minimum_scale = max(0, round(lo_all * 0.8, -1))

    # 5. DCF
    if dcf is not None and not dcf.result.reason:
        res = dcf.result
        s = deck.slide("Discounted cash flow", f"Unlevered DCF, 5-year forecast, perpetuity growth terminal value "
                                               f"(FD-11). WACC {_fmt(res.wacc)}, terminal growth "
                                               f"{_fmt(dcf.assumptions['terminal_growth'])}.")
        rows = [["", *[f"Year {y.n}" for y in res.years]],
                ["Revenue", *[_num(y.revenue, "USD") for y in res.years]],
                ["Growth", *[f"{y.growth:.1%}" for y in res.years]],
                ["EBIT", *[_num(y.ebit, "USD") for y in res.years]],
                ["Unlevered FCF", *[_num(y.ufcf, "USD") for y in res.years]],
                ["PV of FCF", *[_num(y.pv, "USD") for y in res.years]]]
        deck.table(s, rows, Inches(0.5), Inches(1.5), Inches(7.4), Inches(2.4), first_col_width=Inches(1.9), font=11)
        out = [["Output", ""], ["Enterprise value", _fmt(res.enterprise_value)],
               ["Equity value", _fmt(res.equity_value)], ["Share price", _fmt(res.share_price)],
               ["vs current", f"{res.share_price.value / price - 1:+.0%}" if price and res.share_price.ok else "n/a"],
               ["Terminal value % of EV", f"{res.tv_share_of_ev:.0%}" if res.tv_share_of_ev is not None else "n/a"]]
        deck.table(s, out, Inches(8.4), Inches(1.5), Inches(4.4), Inches(2.4), first_col_width=Inches(2.6), font=11)
        ws_, gs, grid = dcf.sensitivity
        if grid:
            sens = [["WACC \\ g", *[f"{g:.1%}" for g in gs]]] + [
                [f"{w:.1%}", *[_fmt(f) for f in row]] for w, row in zip(ws_, grid)]
            deck.table(s, sens, Inches(0.5), Inches(4.3), Inches(7.4), Inches(2.4), first_col_width=Inches(1.4),
                       font=11)
            tb = s.shapes.add_textbox(Inches(8.4), Inches(4.3), Inches(4.4), Inches(2.4)).text_frame
            tb.word_wrap = True
            tb.text = ("Sensitivity: share price by WACC (rows) and terminal growth (columns). History-based "
                       "defaults; every assumption is listed in the Excel export.")
            tb.paragraphs[0].runs[0].font.size = Pt(12)
    return deck.save()


def sector_deck(s: Sector, points: list[str] | None = None) -> bytes:
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches, Pt

    from engine.sector import fit

    rows = s.loaded
    as_of = max((r.get("price").as_of for r in rows if r.get("price").as_of), default=None)
    deck = _Deck(f"Source: SEC EDGAR XBRL filings (each company's latest 10-K/10-Q, LTM); prices from Yahoo Finance "
                 f"as of {as_of}. NM = not meaningful. Statistics exclude NM values (FD-12). Built with Tearsheet.")

    # 1. Overview
    sl = deck.slide(f"{s.name}: sector overview", f"{len(rows)} companies · {date.today():%B %d, %Y}")
    if points:
        deck.bullets(sl, points, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3), size=16)

    # 2. Benchmarking table
    sl = deck.slide(f"{s.name}: trading comparables", "LTM figures; sorted by market cap.")
    keys = ("equity_value", "enterprise_value", "revenue_growth_ltm", "ebitda_margin", "ev_revenue", "ev_ebitda", "pe")
    head = ["Company", "Mkt cap", "EV", "Rev growth", "EBITDA mgn", "EV/Rev", "EV/EBITDA", "P/E"]
    ordered = sorted(rows, key=lambda r: r.get("equity_value").value or 0, reverse=True)
    data = [head] + [[r.name[:28], *[_fmt(r.get(k)) for k in keys]] for r in ordered]
    stat_start = len(data)
    for label, attr in (("Median", "median"), ("25th percentile", "p25"), ("75th percentile", "p75")):
        data.append([label, *[_num(getattr(s.stats[k], attr), SECTOR_UNITS[k]) for k in keys]])
    n = len(data)
    deck.table(sl, data, Inches(0.5), Inches(1.5), Inches(12.3), Inches(min(5.3, 0.3 * n)),
               bold_rows=set(range(stat_start, n)), first_col_width=Inches(3.3), font=11 if n < 16 else 9)

    # 3. Ranked EV/EBITDA and EV/Revenue
    for key in ("ev_ebitda", "ev_revenue"):
        ranked = s.rank(key)
        if not ranked:
            continue
        st = s.stats[key]
        sl = deck.slide(f"{s.name}: LTM {SECTOR_LABELS[key]}",
                        f"Sector median {_num(st.median, 'x')} shown in orange. Left out as NM or n/a: "
                        + (", ".join(t for t, _ in st.excluded) or "none") + ".")
        cd = CategoryChartData()
        cd.categories = [t for t, _ in ranked] + ["Median"]
        cd.add_series(SECTOR_LABELS[key], [round(v, 2) for _, v in ranked] + [round(st.median, 2)])
        ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(12.3),
                                 Inches(5.3), cd).chart
        _style_chart(ch, legend=False)
        ser = ch.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        pt = ser.points[len(ranked)]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(ORANGE)
        plot = ch.plots[0]
        plot.gap_width = 50
        plot.has_data_labels = True
        plot.data_labels.number_format, plot.data_labels.number_format_is_linked = '0.0"x"', False
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False

    # 4. Scatter: EV/Revenue vs growth with the fitted line
    f = fit(rows, "revenue_growth_ltm", "ev_revenue")
    pts = [(r.ticker, r.get("revenue_growth_ltm").value, r.get("ev_revenue").value) for r in rows
           if r.get("revenue_growth_ltm").ok and r.get("ev_revenue").ok]
    if len(pts) >= 3:
        sub = (f"Least-squares line: EV/Revenue = {f.intercept:.1f}x + {f.slope:.1f}x × growth, R² = {f.r2:.2f}, "
               f"n = {f.n}." if f.r2 is not None else "Too few points for a fitted line.")
        sl = deck.slide(f"{s.name}: valuation vs growth", sub)
        cd = XyChartData()
        ser = cd.add_series("Companies")
        for _, x, y in pts:
            ser.add_data_point(round(x, 4), round(y, 2))
        if f.slope is not None:
            xs = [min(p[1] for p in pts), max(p[1] for p in pts)]
            line = cd.add_series("Fit")
            for x in xs:
                line.add_data_point(round(x, 4), round(f.predict(x), 2))
        ch = sl.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3),
                                 cd).chart
        _style_chart(ch, legend=False)
        dots = ch.series[0]
        dots.format.line.fill.background()
        dots.marker.style = 8  # circle
        dots.marker.size = 10
        dots.marker.format.fill.solid()
        dots.marker.format.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        for i, (tk, _, _) in enumerate(pts):
            dl = dots.points[i].data_label
            dl.has_text_frame = True
            dl.text_frame.text = tk
            dl.position = XL_LABEL_POSITION.RIGHT
            dl.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
        if f.slope is not None:
            fl = ch.series[1]
            fl.format.line.color.rgb = RGBColor.from_string(ORANGE)
            fl.format.line.dash_style = 4
            fl.smooth = False
            from pptx.enum.chart import XL_MARKER_STYLE
            fl.marker.style = XL_MARKER_STYLE.NONE
        ch.category_axis.tick_labels.number_format = "0%"
        ch.category_axis.tick_labels.number_format_is_linked = False
        ch.category_axis.has_title = True
        ch.category_axis.axis_title.text_frame.text = "LTM revenue growth"
        ch.value_axis.tick_labels.number_format = '0.0"x"'
        ch.value_axis.tick_labels.number_format_is_linked = False
        ch.value_axis.has_title = True
        ch.value_axis.axis_title.text_frame.text = "EV / LTM revenue"
        ch.value_axis.has_major_gridlines = True
    return deck.save()
