"""Presentation-grade DCF model in Excel (FD-11, FD-13).

The app's DCF (engine/dcf.py) rebuilt as a live, formula-driven workbook in the
layout a bank analyst would hand to a VP:

    Cover        key outputs, valuation summary table and football-field chart
    Inputs       every assumption in one place (blue = typed input), with the
                 historical figure each default came from
    DCF          historicals (10-K fiscal years + LTM), 5-year projection,
                 terminal year, then enterprise value -> share price by both
                 perpetuity growth and exit multiple
    WACC         CAPM cost of equity, after-tax cost of debt, market weights
    Sensitivity  share price by WACC x terminal growth and WACC x exit multiple,
                 and upside vs the current price
    Sources      XBRL tag, filing and formula behind every hard-coded number

Every projection, discounting and valuation cell is an Excel formula that
points back at the Inputs sheet, so changing a blue cell reflows the whole
model. On the app's assumptions the file reproduces the app's numbers exactly;
a "Checks" block on the DCF sheet shows the tie-out.

No finance logic of its own: the formulas are engine/dcf.py written in Excel,
and the only numbers typed in are figures the engine already computed and
sourced (CLAUDE.md rules 1-2).
"""

from __future__ import annotations

import io
import re
from copy import copy
from datetime import date

from engine import dcf as dcf_mod
from engine.export import DcfExport, _Book
from engine.periods import Figure
from engine.report import Report

# ---- house style ------------------------------------------------------------------------

FONT = "Arial"
NAVY, MID, PALE, RULE = "1F3864", "D9E1F2", "F2F5FB", "8EA9DB"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "7F7F7F"
INPUT_FILL = "FFF2CC"  # pale yellow: cells meant to be changed

FMT = {
    "usd": '#,##0_);(#,##0);"–"_)',
    "usd1": '#,##0.0_);(#,##0.0);"–"_)',
    "px": '$#,##0.00_);($#,##0.00);"–"_)',
    "pct": '0.0%_);(0.0%);"–"_)',
    "pct2": '0.00%_);(0.00%);"–"_)',
    "x": '0.0"x"_);(0.0"x");"–"_)',
    "beta": '0.00_);(0.00)',
    "factor": '0.000_)',
    "period": '0.0_)',
    "flag": '0',
    "shares": '#,##0_)',
}
UNIT_FMT = {"USD": "usd", "shares": "shares", "USD/share": "px", "x": "x", "ratio": "pct"}

# Sensitivity grids (FD-11 steps, as in the app)
EXIT_STEPS = (-2.0, -1.0, 0.0, 1.0, 2.0)


_PURE_LINK = re.compile(r"^=-?'?[A-Za-z ]+'?!\$?[A-Z]+\$?\d+$")  # a bare link to another sheet


def _m(f: Figure | None):
    """Figure -> sheet number ($ and shares in millions), or "n/a"/"NM"."""
    if f is None or f.value is None:
        return "NM" if f is not None and f.nm else "n/a"
    return f.value / 1e6 if f.unit in ("USD", "shares") else f.value


def _plus_years(d: date, n: int) -> date:
    try:
        return d.replace(year=d.year + n)
    except ValueError:  # 29 Feb
        return d.replace(year=d.year + n, day=28)


class _Sheet:
    """openpyxl worksheet with banker styling helpers."""

    def __init__(self, ws, book: "_Model"):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self.ws, self.book = ws, book
        self.Font, self.Fill, self.Align, self.Border, self.Side = Font, PatternFill, Alignment, Border, Side
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.column_dimensions["A"].width = 2

    # cells
    def put(self, ref: str, value, fmt: str | None = None, *, color: str = BLACK, bold=False, italic=False,
            fill: str | None = None, align: str | None = None, size: int = 10, wrap=False, indent: int = 0):
        c = self.ws[ref]
        c.value = value
        c.font = self.Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
        if fmt:
            c.number_format = FMT.get(fmt, fmt)
        if fill:
            c.fill = self.Fill("solid", fgColor=fill)
        if align or wrap or indent:
            c.alignment = self.Align(horizontal=align, wrap_text=wrap, indent=indent, vertical="center")
        elif isinstance(value, str) and not value.startswith("=") and value in ("n/a", "NM"):
            c.alignment = self.Align(horizontal="right")
        return c

    def inp(self, ref: str, value, fmt: str, fill: bool = True):
        """Typed-in number: blue font (pale yellow fill when it is an assumption to play with)."""
        return self.put(ref, value, fmt, color=BLUE, fill=INPUT_FILL if fill else None)

    def link(self, ref: str, formula: str, fmt: str | None = None, **kw):
        """Formula pulling from another sheet: green."""
        return self.put(ref, formula, fmt, color=GREEN, **kw)

    # blocks
    def title(self, heading: str, sub: str, last_col: str):
        ws = self.ws
        for col in range(2, ws[f"{last_col}1"].column + 1):
            ws.cell(row=1, column=col).fill = self.Fill("solid", fgColor=NAVY)
            ws.cell(row=2, column=col).fill = self.Fill("solid", fgColor=NAVY)
        self.put("B1", heading, color="FFFFFF", bold=True, size=14)
        self.put("B2", sub, color="FFFFFF", italic=True, size=9)
        ws.row_dimensions[1].height = 24

    def band(self, row: int, first: str, last: str, text: str = "", values: dict | None = None):
        """Navy column-header band across first..last."""
        ws = self.ws
        a, b = ws[f"{first}{row}"].column, ws[f"{last}{row}"].column
        for col in range(a, b + 1):
            c = ws.cell(row=row, column=col)
            c.fill = self.Fill("solid", fgColor=NAVY)
            c.font = self.Font(name=FONT, size=10, bold=True, color="FFFFFF")
            c.alignment = self.Align(horizontal="center" if col > a else "left", vertical="center", wrap_text=col > a)
        ws.cell(row=row, column=a, value=text)
        for ref, v in (values or {}).items():
            ws[f"{ref}{row}"].value = v

    def section(self, row: int, first: str, last: str, text: str):
        """Bold label with a light rule under it."""
        ws = self.ws
        a, b = ws[f"{first}{row}"].column, ws[f"{last}{row}"].column
        for col in range(a, b + 1):
            c = ws.cell(row=row, column=col)
            c.fill = self.Fill("solid", fgColor=MID)
            c.font = self.Font(name=FONT, size=10, bold=True, color=NAVY)
        ws.cell(row=row, column=a, value=text)

    def total(self, row: int, first: str, last: str, double: bool = False):
        """Bold the row and put a rule above it (double rule for the answer)."""
        ws = self.ws
        a, b = ws[f"{first}{row}"].column, ws[f"{last}{row}"].column
        top = self.Side(style="thin", color=BLACK)
        bottom = self.Side(style="double", color=BLACK) if double else None
        for col in range(a, b + 1):
            c = ws.cell(row=row, column=col)
            f = copy(c.font)
            f.bold = True
            c.font = f
            c.border = self.Border(top=top, bottom=bottom) if bottom else self.Border(top=top)

    def shade(self, ref_range: str, color: str = PALE):
        for row in self.ws[ref_range]:
            for c in row:
                c.fill = self.Fill("solid", fgColor=color)

    def widths(self, **w):
        for col, width in w.items():
            self.ws.column_dimensions[col].width = width

    def landscape(self, area: str):
        ws = self.ws
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = area
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4


class _Model:
    def __init__(self):
        self.book = _Book()  # reuse the Sources sheet and save logic
        self.wb = self.book.wb

    def sheet(self, name: str) -> _Sheet:
        return _Sheet(self.wb.create_sheet(name), self)


# ---- the workbook -----------------------------------------------------------------------


def dcf_workbook(rep: Report, d: DcfExport) -> bytes:
    """The full DCF model for the report's target, on the app's current assumptions."""
    t, b, res, a = rep.target, rep.dcf_base, d.result, d.assumptions
    if res.reason:
        raise ValueError(f"DCF not available: {res.reason}")
    m, F = t.metrics, t.metrics.figures
    model = _Model()
    log = model.book.log
    cik = t.cik
    hist = list(reversed(b.years))  # oldest first
    H = len(hist)
    price_date = F["price"].as_of or date.today()
    name = f"{t.name} ({t.ticker})"
    unit_line = "$ in millions, except per-share data"

    cover = model.sheet("Cover")
    inputs = model.sheet("Inputs")
    dcf = model.sheet("DCF")
    wacc = model.sheet("WACC")
    sens = model.sheet("Sensitivity")

    # ======================= Inputs =======================
    I = {}  # key -> absolute reference "Inputs!$C$n"
    s = inputs
    s.title(f"{name}: DCF inputs", f"{unit_line}. Blue cells on yellow are inputs: change them and the whole "
                                   "model updates. Black = formula, green = link to another sheet.", "E")
    s.widths(B=44, C=14, D=14, E=96)

    def iref(key: str, row: int) -> None:
        I[key] = f"Inputs!$C${row}"

    r = 4
    s.band(r, "B", "E", "Market data and balance sheet", {"C": "Value", "D": "", "E": "Source"})
    as_of_bs = F["total_debt"].as_of or m.ltm_end
    market_rows = [
        ("price", "Current share price ($)", F["price"], "px"),
        ("shares", "Shares outstanding (m, latest 10-K/10-Q cover page)", F["shares_outstanding"], "shares"),
        ("debt", f"Total debt incl. finance leases (at {as_of_bs})", F["total_debt"], "usd"),
        ("preferred", "Preferred stock", F["preferred_stock"], "usd"),
        ("nci", "Noncontrolling interest", F["noncontrolling_interest"], "usd"),
        ("cash", "Cash and equivalents", F["cash"], "usd"),
        ("sti", "Short-term investments", F["short_term_investments"], "usd"),
        ("ltm_rev", f"LTM revenue (12 months to {m.ltm_end})", F["ltm_revenue"], "usd"),
    ]
    for key, label, f, fmt in market_rows:
        r += 1
        s.put(f"B{r}", label)
        s.inp(f"C{r}", _m(f), fmt, fill=False)
        s.put(f"E{r}", f.formula if f.ok else f"n/a: {f.reason}", color=GREY, size=9)
        iref(key, r)
        log("Inputs", label, f, cik)
    r += 1
    s.put(f"B{r}", "Market capitalization (price × shares)")
    s.put(f"C{r}", f"={I['price']}*{I['shares']}".replace("Inputs!", ""), "usd")
    s.put(f"E{r}", "Equity value at the current price (FD-3)", color=GREY, size=9)
    iref("mcap", r)

    r += 2
    s.band(r, "B", "E", "Revenue growth", {"C": "Assumption", "D": "Historical", "E": "Basis"})
    for i, k in enumerate(dcf_mod.GROWTH_KEYS, 1):
        r += 1
        f = a[k]
        s.put(f"B{r}", f"Year {i}  ({_plus_years(m.ltm_end, i - 1):%b-%y} to {_plus_years(m.ltm_end, i):%b-%y})")
        s.inp(f"C{r}", f.value, "pct")
        s.put(f"E{r}", f.formula, color=GREY, size=9)
        iref(k, r)
        log("Assumptions", dcf_mod.LABELS[k], f, cik)
    I_growth_hist_row = r - dcf_mod.YEARS + 1  # filled once the DCF sheet exists

    r += 2
    s.band(r, "B", "E", "Operating assumptions (% of revenue unless noted)",
           {"C": "Assumption", "D": "Historical", "E": "Basis"})
    op_rows = {}
    for k in dcf_mod.OPERATING_KEYS:
        r += 1
        f = a[k]
        s.put(f"B{r}", dcf_mod.LABELS[k] + (" (on positive EBIT)" if k == "tax_rate" else ""))
        s.inp(f"C{r}", f.value, "pct")
        s.put(f"E{r}", f.formula, color=GREY, size=9)
        iref(k, r)
        op_rows[k] = r
        log("Assumptions", dcf_mod.LABELS[k], f, cik)

    r += 2
    s.band(r, "B", "E", "Cost of capital", {"C": "Assumption", "D": "", "E": "Basis"})
    for k, label, fmt in (("risk_free", "Risk-free rate (10-year US Treasury)", "pct2"),
                          ("erp", "Equity risk premium", "pct2")):
        r += 1
        s.put(f"B{r}", label)
        s.inp(f"C{r}", a[k].value, fmt)
        s.put(f"E{r}", a[k].formula, color=GREY, size=9)
        iref(k, r)
        log("Assumptions", dcf_mod.LABELS[k], a[k], cik)
    raw = b.beta
    r += 1
    s.put(f"B{r}", "Raw beta (5-year monthly vs S&P 500)")
    s.inp(f"C{r}", raw.value if raw.ok else "n/a", "beta", fill=False)
    s.put(f"E{r}", raw.formula if raw.ok else f"n/a: {raw.reason}", color=GREY, size=9)
    iref("raw_beta", r)
    log("Assumptions", "Raw beta", raw, cik)
    r += 1
    s.put(f"B{r}", "Weight on raw beta (Blume adjustment)")
    s.inp(f"C{r}", dcf_mod.BETA_RAW_WEIGHT, "beta")
    s.put(f"E{r}", "Adjusted beta = weight × raw + (1 − weight) × 1.0: betas drift toward the market's 1.0 "
                   "(Bloomberg convention, FD-11)", color=GREY, size=9)
    iref("beta_w", r)
    r += 1
    s.put(f"B{r}", "Adjusted beta")
    beta = a["beta"]
    if beta.formula.startswith("entered by user") or not raw.ok:
        s.inp(f"C{r}", beta.value, "beta")
        s.put(f"E{r}", f"Beta {beta.formula} in the app; typed in here", color=GREY, size=9)
    else:
        s.put(f"C{r}", f"=C{r - 1}*C{r - 2}+(1-C{r - 1})", "beta")
        s.put(f"E{r}", "Formula: change the raw beta or the weight above, or overwrite this cell", color=GREY,
              size=9)
    iref("beta", r)
    r += 1
    s.put(f"B{r}", "Pre-tax cost of debt")
    s.inp(f"C{r}", a["cost_of_debt"].value, "pct2")
    s.put(f"E{r}", a["cost_of_debt"].formula, color=GREY, size=9)
    iref("cost_of_debt", r)
    log("Assumptions", dcf_mod.LABELS["cost_of_debt"], a["cost_of_debt"], cik)

    r += 2
    s.band(r, "B", "E", "Terminal value and conventions", {"C": "Assumption", "D": "", "E": "Basis"})
    r += 1
    s.put(f"B{r}", "Terminal growth rate (perpetuity method)")
    s.inp(f"C{r}", a["terminal_growth"].value, "pct2")
    s.put(f"E{r}", a["terminal_growth"].formula, color=GREY, size=9)
    iref("g", r)
    log("Assumptions", dcf_mod.LABELS["terminal_growth"], a["terminal_growth"], cik)
    r += 1
    peer = b.peer_median_ev_ebitda
    s.put(f"B{r}", "Exit EV / EBITDA multiple (exit-multiple method)")
    exit_row = r
    iref("exit", r)
    if peer is not None and peer.ok:
        s.inp(f"C{r}", peer.value, "x")
        s.put(f"E{r}", f"EV / LTM EBITDA, {peer.formula}, from the Tearsheet comps", color=GREY, size=9)
        log("Assumptions", "Exit EV/EBITDA multiple", peer, cik)
    r += 1
    s.put(f"B{r}", "Mid-year convention (1 = on, 0 = off)")
    s.inp(f"C{r}", 1 if d.mid_year else 0, "flag")
    s.put(f"E{r}", "On: each year's cash arrives mid-year, discounted half a year less (Rosenbaum & Pearl). "
                   "The exit-multiple value is always discounted from the end of year 5.", color=GREY, size=9)
    iref("mid", r)
    r += 1
    s.put(f"B{r}", "Steady-state terminal year (1 = on, 0 = off)")
    s.inp(f"C{r}", 1 if d.steady_state else 0, "flag")
    s.put(f"E{r}", "On: terminal-year capex is capped at D&A, since capex above depreciation is growth "
                   "investment a company growing at the terminal rate doesn't keep making (FD-11)", color=GREY,
          size=9)
    iref("steady", r)

    r += 2
    s.band(r, "B", "E", "Sensitivity table steps", {"C": "Step", "D": "", "E": ""})
    for key, label, v, fmt in (("w_step", "WACC step", 0.005, "pct2"), ("g_step", "Terminal growth step", 0.005,
                                                                         "pct2"),
                               ("x_step", "Exit multiple step", 1.0, "x")):
        r += 1
        s.put(f"B{r}", label)
        s.inp(f"C{r}", v, fmt)
        iref(key, r)

    r += 2
    s.band(r, "B", "E", "Tearsheet app output on these assumptions (for the tie-out check)",
           {"C": "Value", "D": "", "E": ""})
    for key, label, f, fmt in (("app_price", "DCF share price (perpetuity growth)", res.share_price, "px"),
                               ("app_ev", "DCF enterprise value", res.enterprise_value, "usd"),
                               ("app_wacc", "WACC", res.wacc, "pct2")):
        r += 1
        s.put(f"B{r}", label)
        s.put(f"C{r}", _m(f), fmt, color=GREY)
        s.put(f"E{r}", "Hard-coded from the app when this file was made. The Checks block on the DCF sheet "
                       "compares the model with it.", color=GREY, size=9)
        iref(key, r)
    inputs.ws.freeze_panes = "C4"
    inputs.landscape(f"A1:E{r}")

    # ======================= DCF sheet =======================
    s = dcf
    hc = [chr(ord("C") + i) for i in range(H)]  # historical fiscal years
    lc = chr(ord("C") + H)  # LTM
    pc = [chr(ord(lc) + i) for i in range(1, dcf_mod.YEARS + 1)]  # projection years
    tc = chr(ord(pc[-1]) + 1)  # terminal column
    first_p, last_p = pc[0], pc[-1]
    s.title(f"{name}: discounted cash flow analysis",
            f"{unit_line}. Historicals from 10-K filings (fiscal years) and 10-K/10-Q (LTM to {m.ltm_end}); "
            f"projections are formulas on the Inputs sheet.", tc)
    s.widths(B=40)
    for c in hc + [lc, *pc, tc]:
        s.ws.column_dimensions[c].width = 12.5
    R = {}
    r = 4
    s.band(r, "B", tc, "", {**{c: f"FY{y.end.year}A" for c, y in zip(hc, hist)}, lc: "LTM",
                            **{c: f"Year {i}E" for i, c in enumerate(pc, 1)}, tc: "Terminal"})
    r += 1
    s.put(f"B{r}", "Period end", italic=True, color=GREY, size=9)
    for c, y in zip(hc, hist):
        s.put(f"{c}{r}", f"{y.end:%b-%y}", italic=True, color=GREY, size=9, align="center")
    s.put(f"{lc}{r}", f"{m.ltm_end:%b-%y}", italic=True, color=GREY, size=9, align="center")
    for i, c in enumerate(pc, 1):
        s.put(f"{c}{r}", f"{_plus_years(m.ltm_end, i):%b-%y}", italic=True, color=GREY, size=9, align="center")
    s.put(f"{tc}{r}", "normalized yr 5", italic=True, color=GREY, size=9, align="center")
    for c in [*hc, lc, *pc, tc]:
        s.ws[f"{c}4"].alignment = s.Align(horizontal="center")

    def row(key: str, label: str, *, bold=False, indent=0, italic=False):
        nonlocal r
        r += 1
        R[key] = r
        s.put(f"B{r}", label, bold=bold, italic=italic, indent=indent, color=GREY if italic else BLACK)
        return r

    def hist_vals(rw: int, attr: str, label: str, sign: int = 1):
        for c, y in zip(hc, hist):
            f = getattr(y, attr)
            v = _m(f)
            s.inp(f"{c}{rw}", sign * v if isinstance(v, float) else v, "usd", fill=False)
            log("DCF historicals", f"{label}, FY{y.end.year}", f, cik)

    def pct_row(key: str, label: str, num: str, cols, sign: str = ""):
        rw = row(key, label, italic=True, indent=1)
        for c in cols:
            s.put(f"{c}{rw}", f'=IFERROR({sign}{c}{R[num]}/{c}{R["rev"]},"n/a")', "pct", italic=True, color=GREY)

    # Revenue
    rw = row("rev", "Revenue", bold=True)
    hist_vals(rw, "revenue", "Revenue")
    s.link(f"{lc}{rw}", f"={I['ltm_rev']}", "usd", bold=True)
    prev = lc
    for i, c in enumerate(pc):
        s.put(f"{c}{rw}", f"={prev}{rw}*(1+{c}{rw + 1})", "usd", bold=True)
        prev = c
    s.put(f"{tc}{rw}", f"={last_p}{rw}", "usd", bold=True)
    rw = row("growth", "% growth", italic=True, indent=1)
    for j, c in enumerate(hc):
        if j > 0:
            s.put(f"{c}{rw}", f'=IFERROR({c}{rw - 1}/{hc[j - 1]}{rw - 1}-1,"n/a")', "pct", italic=True, color=GREY)
    g_ltm = F["revenue_growth_ltm"]
    s.inp(f"{lc}{rw}", _m(g_ltm), "pct", fill=False)
    log("DCF historicals", "Revenue growth, LTM", g_ltm, cik)
    for i, c in enumerate(pc):
        s.link(f"{c}{rw}", f"={I[dcf_mod.GROWTH_KEYS[i]]}", "pct", italic=True)

    r += 1
    rw = row("ebit", "EBIT (operating income)")
    hist_vals(rw, "ebit", "EBIT")
    s.inp(f"{lc}{rw}", _m(F["ltm_operating_income"]), "usd", fill=False)
    log("DCF historicals", "EBIT, LTM", F["ltm_operating_income"], cik)
    for c in pc:
        s.put(f"{c}{rw}", f"={c}{R['rev']}*{I['ebit_margin']}", "usd")
    s.put(f"{tc}{rw}", f"={last_p}{rw}", "usd")
    pct_row("ebit_m", "% margin", "ebit", [*hc, lc, *pc, tc])
    rw = row("tax", "(−) Taxes on EBIT")
    for c in [*pc, tc]:
        s.put(f"{c}{rw}", f"=-MAX({c}{R['ebit']},0)*{I['tax_rate']}", "usd")
    rw = row("nopat", "NOPAT", bold=True)
    for c in [*pc, tc]:
        s.put(f"{c}{rw}", f"={c}{R['ebit']}+{c}{R['tax']}", "usd")
    s.total(rw, first_p, tc)
    rw = row("da", "(+) Depreciation & amortization")
    hist_vals(rw, "d_and_a", "D&A")
    s.inp(f"{lc}{rw}", _m(F["ltm_d_and_a"]), "usd", fill=False)
    log("DCF historicals", "D&A, LTM", F["ltm_d_and_a"], cik)
    for c in pc:
        s.put(f"{c}{rw}", f"={c}{R['rev']}*{I['da_pct']}", "usd")
    s.put(f"{tc}{rw}", f"={last_p}{rw}", "usd")
    pct_row("da_p", "% of revenue", "da", [*hc, lc, *pc, tc])
    rw = row("capex", "(−) Capital expenditures")
    hist_vals(rw, "capex", "Capex", sign=-1)
    for c in pc:
        s.put(f"{c}{rw}", f"=-{c}{R['rev']}*{I['capex_pct']}", "usd")
    # FD-11 steady state: capex no bigger than D&A in the perpetuity (both negative here, so MAX).
    s.put(f"{tc}{rw}", f"=IF({I['steady']}=1,MAX({last_p}{rw},-{tc}{R['da']}),{last_p}{rw})", "usd")
    pct_row("capex_p", "% of revenue", "capex", [*hc, *pc, tc], sign="-")
    rw = row("dnwc", "(−) Increase in net working capital")
    prev = lc
    for c in pc:
        s.put(f"{c}{rw}", f"=-({c}{R['rev']}-{prev}{R['rev']})*{I['nwc_pct']}", "usd")
        prev = c
    s.put(f"{tc}{rw}", f"={last_p}{rw}", "usd")
    rw = row("ufcf", "Unlevered free cash flow", bold=True)
    for c in [*pc, tc]:
        s.put(f"{c}{rw}", f"={c}{R['nopat']}+{c}{R['da']}+{c}{R['capex']}+{c}{R['dnwc']}", "usd")
    s.total(rw, "B", tc, double=True)
    pct_row("ufcf_m", "% margin", "ufcf", [*pc])

    r += 1
    rw = row("period", "Discount period (years)")
    for i, c in enumerate(pc, 1):
        s.put(f"{c}{rw}", f"={i}-0.5*{I['mid']}", "period")
    rw = row("df", "Discount factor at WACC")
    for c in pc:
        s.put(f"{c}{rw}", f"=1/(1+WACC!$D$WACCROW)^{c}{R['period']}", "factor")
    rw = row("pv", "Present value of free cash flow", bold=True)
    for c in pc:
        s.put(f"{c}{rw}", f"={c}{R['ufcf']}*{c}{R['df']}", "usd", bold=True)

    r += 1
    rw = row("memo", "Memo", bold=True)
    s.section(rw, "B", tc, "Memo: historical drivers")
    rw = row("ebitda", "EBITDA (EBIT + D&A)")
    for c in [*hc, lc, *pc]:
        s.put(f"{c}{rw}", f'=IFERROR({c}{R["ebit"]}+{c}{R["da"]},"n/a")', "usd")
    rw = row("nwc", "Net working capital (operating)")
    hist_vals(rw, "nwc", "Net working capital")
    for c in pc:
        s.put(f"{c}{rw}", f"={c}{R['rev']}*{I['nwc_pct']}", "usd")
    pct_row("nwc_p", "% of revenue", "nwc", [*hc, *pc])
    rw = row("inc_tax", "Income tax expense")
    hist_vals(rw, "income_tax", "Income tax expense")
    rw = row("pretax", "Pretax income")
    hist_vals(rw, "pretax_income", "Pretax income")
    rw = row("etr", "Effective tax rate", italic=True, indent=1)
    for c in hc:
        s.put(f"{c}{rw}", f'=IFERROR({c}{R["inc_tax"]}/{c}{R["pretax"]},"n/a")', "pct", italic=True, color=GREY)
    # Lightly shade the LTM and terminal columns so the history / forecast split reads at a glance.
    s.shade(f"{lc}6:{lc}{r}", PALE)
    s.shade(f"{tc}6:{tc}{R['ufcf_m']}", PALE)

    # Valuation: perpetuity growth vs exit multiple, side by side
    r += 2
    V = {}
    s.band(r, "B", "D", "Valuation", {"C": "Perpetuity growth", "D": "Exit multiple"})
    s.ws.row_dimensions[r].height = 28

    def vrow(key: str, label: str, perp, exit_, fmt: str, *, bold=False, color=BLACK, total=False, double=False):
        nonlocal r
        r += 1
        V[key] = r
        s.put(f"B{r}", label, bold=bold)
        for col, v in (("C", perp), ("D", exit_)):
            if v is None:
                continue
            green = isinstance(v, str) and _PURE_LINK.match(v) is not None
            s.put(f"{col}{r}", v, fmt, bold=bold, color=GREEN if green and color == BLACK else color)
        if total or double:
            s.total(r, "B", "D", double=double)

    hc_n = {"ufcf": f"{tc}{R['ufcf']}", "ebitda5": f"{last_p}{R['ebitda']}"}
    vrow("metric", "Terminal-year UFCF / Year-5 EBITDA", f"={hc_n['ufcf']}", f"={hc_n['ebitda5']}", "usd")
    vrow("driver", "Terminal growth / exit EV/EBITDA", f"={I['g']}", f"={I['exit']}", "pct2")
    s.ws[f"D{V['driver']}"].number_format = FMT["x"]
    vrow("tv", "Terminal value",
         f'=IF(OR(WACC!$D$WACCROW<=C{r},C{r - 1}<=0),"NM",C{r - 1}*(1+C{r})/(WACC!$D$WACCROW-C{r}))',
         f"=D{r - 1}*D{r}", "usd")
    vrow("tvp", "Discount period for terminal value", f"={dcf_mod.YEARS}-0.5*{I['mid']}", f"={dcf_mod.YEARS}",
         "period")
    vrow("pv_tv", "PV of terminal value", f'=IFERROR(C{V["tv"]}/(1+WACC!$D$WACCROW)^C{V["tvp"]},"NM")',
         f'=IFERROR(D{V["tv"]}/(1+WACC!$D$WACCROW)^D{V["tvp"]},"NM")', "usd")
    vrow("pv_fcf", "PV of free cash flows (years 1-5)", f"=SUM({first_p}{R['pv']}:{last_p}{R['pv']})",
         f"=SUM({first_p}{R['pv']}:{last_p}{R['pv']})", "usd")
    vrow("ev", "Enterprise value", f'=IFERROR(C{V["pv_fcf"]}+C{V["pv_tv"]},"NM")',
         f'=IFERROR(D{V["pv_fcf"]}+D{V["pv_tv"]},"NM")', "usd", bold=True, total=True)
    bridge_first = r + 1
    for key, label, ref, sign in (("debt", "(−) Total debt", I["debt"], "-"),
                                  ("preferred", "(−) Preferred stock", I["preferred"], "-"),
                                  ("nci", "(−) Noncontrolling interest", I["nci"], "-"),
                                  ("cash", "(+) Cash and equivalents", I["cash"], ""),
                                  ("sti", "(+) Short-term investments", I["sti"], "")):
        vrow(key, label, f"={sign}{ref}", f"={sign}{ref}", "usd", color=GREEN)
    bridge_last = r
    vrow("eq", "Equity value", f'=IFERROR(C{V["ev"]}+SUM(C{bridge_first}:C{bridge_last}),"NM")',
         f'=IFERROR(D{V["ev"]}+SUM(D{bridge_first}:D{bridge_last}),"NM")', "usd", bold=True, total=True)
    vrow("shares", "(÷) Shares outstanding (m)", f"={I['shares']}", f"={I['shares']}", "shares", color=GREEN)
    vrow("px", "Implied share price", f'=IFERROR(IF(C{V["eq"]}<=0,"NM",C{V["eq"]}/C{V["shares"]}),"NM")',
         f'=IFERROR(IF(D{V["eq"]}<=0,"NM",D{V["eq"]}/D{V["shares"]}),"NM")', "px", bold=True, double=True)
    vrow("cur", "Current share price", f"={I['price']}", f"={I['price']}", "px", color=GREEN)
    vrow("up", "Upside / (downside)", f'=IFERROR(C{V["px"]}/C{V["cur"]}-1,"NM")',
         f'=IFERROR(D{V["px"]}/D{V["cur"]}-1,"NM")', "pct", bold=True)
    r += 1
    vrow("tv_share", "Terminal value as % of EV", f'=IFERROR(C{V["pv_tv"]}/C{V["ev"]},"NM")',
         f'=IFERROR(D{V["pv_tv"]}/D{V["ev"]},"NM")', "pct")
    vrow("implied", "Implied exit multiple / implied perpetual growth",
         f'=IFERROR(IF({hc_n["ebitda5"]}<=0,"NM",C{V["tv"]}/{hc_n["ebitda5"]}),"NM")',
         f'=IFERROR((D{V["tv"]}*WACC!$D$WACCROW-{hc_n["ufcf"]})/(D{V["tv"]}+{hc_n["ufcf"]}),"NM")', "x")
    s.ws[f"D{V['implied']}"].number_format = FMT["pct2"]
    for k in ("px", "up"):
        for col in "CD":
            s.ws[f"{col}{V[k]}"].fill = s.Fill("solid", fgColor=MID)

    if peer is None or not peer.ok:
        # No peer median: start the exit multiple at the one the perpetuity method implies, flagged.
        inputs.put(f"C{exit_row}", f"=DCF!C{V['implied']}", "x", color=GREEN, fill=INPUT_FILL)
        inputs.put(f"E{exit_row}", "No peer median EV/EBITDA was available, so this starts at the multiple the "
                                   "perpetuity method implies (both methods then agree). Type a multiple over it.",
                   color=GREY, size=9)

    # Checks
    r += 2
    s.section(r, "B", "D", "Checks")
    r += 1
    s.put(f"B{r}", "Share price in the Tearsheet app (perpetuity)")
    s.link(f"C{r}", f"={I['app_price']}", "px")
    r += 1
    s.put(f"B{r}", "Difference vs this model")
    s.put(f"C{r}", f'=IFERROR(C{V["px"]}-C{r - 1},"NM")', "px")
    r += 1
    s.put(f"B{r}", "Status")
    s.put(f"C{r}", f'=IF(ISNUMBER(C{r - 1}),IF(ABS(C{r - 1})<0.005,"Ties to the app",'
                   f'"Inputs changed since export"),"NM")', bold=True)
    check_row = r
    r += 1
    s.put(f"B{r}", "WACC above terminal growth")
    s.put(f"C{r}", f'=IF(WACC!$D$WACCROW>{I["g"]},"OK","Fix: WACC must exceed g")')
    notes = list(dict.fromkeys([*res.notes, *(n for f in a.values() for n in f.notes)]))
    if notes:
        r += 2
        s.section(r, "B", tc, "Flags from the app (assumptions worth a second look)")
        for n in notes:
            r += 1
            s.put(f"B{r}", f"• {n}", size=9, color=GREY)
    dcf.ws.freeze_panes = "C6"
    dcf.landscape(f"A1:{tc}{r}")

    # ======================= WACC =======================
    s = wacc
    s.title(f"{name}: weighted average cost of capital", "Market-value weights: equity at today's price, debt "
            "at book (FD-11). Green = linked from Inputs.", "E")
    s.widths(B=4, C=40, D=14, E=70)
    r = 4
    s.band(r, "C", "E", "Cost of equity (CAPM)", {"D": "", "E": ""})
    W = {}
    for key, label, ref, fmt, note in (
            ("rf", "Risk-free rate", I["risk_free"], "pct2", "10-year Treasury yield"),
            ("beta", "Adjusted beta", I["beta"], "beta", "Raw regression beta pulled toward 1.0"),
            ("erp", "Equity risk premium", I["erp"], "pct2", "Expected market return above the risk-free rate")):
        r += 1
        W[key] = r
        s.put(f"C{r}", label)
        s.link(f"D{r}", f"={ref}", fmt)
        s.put(f"E{r}", note, color=GREY, size=9)
    r += 1
    W["ke"] = r
    s.put(f"C{r}", "Cost of equity")
    s.put(f"D{r}", f"=D{W['rf']}+D{W['beta']}*D{W['erp']}", "pct2")
    s.put(f"E{r}", "Risk-free + beta × ERP", color=GREY, size=9)
    s.total(r, "C", "D")

    r += 2
    s.band(r, "C", "E", "Cost of debt", {"D": "", "E": ""})
    for key, label, ref, fmt, note in (
            ("kd", "Pre-tax cost of debt", I["cost_of_debt"], "pct2", "LTM interest expense / total debt"),
            ("t", "Tax rate", I["tax_rate"], "pct", "Interest is tax-deductible")):
        r += 1
        W[key] = r
        s.put(f"C{r}", label)
        s.link(f"D{r}", f"={ref}", fmt)
        s.put(f"E{r}", note, color=GREY, size=9)
    r += 1
    W["kd_at"] = r
    s.put(f"C{r}", "After-tax cost of debt")
    s.put(f"D{r}", f"=D{W['kd']}*(1-D{W['t']})", "pct2")
    s.total(r, "C", "D")

    r += 2
    s.band(r, "C", "E", "Capital structure ($m)", {"D": "", "E": ""})
    r += 1
    W["E"] = r
    s.put(f"C{r}", "Market capitalization")
    s.link(f"D{r}", f"={I['mcap']}", "usd")
    r += 1
    W["D"] = r
    s.put(f"C{r}", "Total debt")
    s.link(f"D{r}", f"={I['debt']}", "usd")
    r += 1
    W["cap"] = r
    s.put(f"C{r}", "Total capitalization")
    s.put(f"D{r}", f"=D{W['E']}+D{W['D']}", "usd")
    s.total(r, "C", "D")
    r += 1
    W["we"] = r
    s.put(f"C{r}", "Equity / total capitalization")
    s.put(f"D{r}", f"=D{W['E']}/D{W['cap']}", "pct")
    r += 1
    W["wd"] = r
    s.put(f"C{r}", "Debt / total capitalization")
    s.put(f"D{r}", f"=D{W['D']}/D{W['cap']}", "pct")
    r += 1
    s.put(f"C{r}", "Debt / equity", italic=True, color=GREY)
    s.put(f"D{r}", f"=D{W['D']}/D{W['E']}", "pct", italic=True, color=GREY)

    r += 2
    W["wacc"] = r
    s.put(f"C{r}", "WACC", bold=True, size=11)
    s.put(f"D{r}", f"=D{W['we']}*D{W['ke']}+D{W['wd']}*D{W['kd_at']}", "pct2", bold=True, size=11, fill=MID)
    s.put(f"E{r}", "Equity weight × cost of equity + debt weight × after-tax cost of debt", color=GREY, size=9)
    s.total(r, "C", "D", double=True)
    wacc.landscape(f"A1:E{r}")
    wacc_ref = f"WACC!$D${W['wacc']}"

    # Patch the WACC placeholder now that its row is known.
    for ws_ in (dcf.ws,):
        for rowc in ws_.iter_rows():
            for c in rowc:
                if isinstance(c.value, str) and "WACCROW" in c.value:
                    c.value = c.value.replace("WACC!$D$WACCROW", wacc_ref)

    # Inputs: "historical" column, linked to the DCF sheet (how each default was derived)
    s = inputs
    hs, he = hc[0], hc[-1]
    hist_formulas = {
        "ebit_margin": f'=IFERROR(SUM(DCF!{hs}{R["ebit"]}:{he}{R["ebit"]})/SUM(DCF!{hs}{R["rev"]}:{he}{R["rev"]}),"n/a")',
        "da_pct": f'=IFERROR(SUM(DCF!{hs}{R["da"]}:{he}{R["da"]})/SUM(DCF!{hs}{R["rev"]}:{he}{R["rev"]}),"n/a")',
        "capex_pct": f'=IFERROR(-SUM(DCF!{hs}{R["capex"]}:{he}{R["capex"]})/SUM(DCF!{hs}{R["rev"]}:{he}{R["rev"]}),"n/a")',
        "nwc_pct": f'=IFERROR(AVERAGE(DCF!{hs}{R["nwc_p"]}:{he}{R["nwc_p"]}),"n/a")',
        "tax_rate": f'=IFERROR(SUM(DCF!{hs}{R["inc_tax"]}:{he}{R["inc_tax"]})/SUM(DCF!{hs}{R["pretax"]}:{he}{R["pretax"]}),"n/a")',
    }
    for k, f in hist_formulas.items():
        s.link(f"D{op_rows[k]}", f, "pct", italic=True)
    s.link(f"D{I_growth_hist_row}", f"=DCF!{lc}{R['growth']}", "pct", italic=True)

    # ======================= Sensitivity =======================
    s = sens
    s.title(f"{name}: sensitivity analysis", "Each cell re-runs the DCF. Green: at or above the current price; red: "
            "below. Steps are inputs.", "H")
    s.widths(B=24, **{c: 13 for c in "CDEFGH"})
    ufcf = f"DCF!${first_p}${R['ufcf']}:${last_p}${R['ufcf']}"
    per = f"DCF!${first_p}${R['period']}:${last_p}${R['period']}"
    tfcf = f"DCF!${tc}${R['ufcf']}"
    ebitda5 = f"DCF!${last_p}${R['ebitda']}"
    claims = f"SUM(DCF!$C${bridge_first}:$C${bridge_last})"
    shares = I["shares"]
    tvp_perp = f"DCF!$C${V['tvp']}"

    def grid(top: int, title: str, col_label: str, col_ref: str, col_step: str, col_fmt: str, cell) -> tuple:
        s.band(top, "B", "H", title)
        s.put(f"C{top + 1}", col_label, italic=True, color=GREY, size=9)
        s.put(f"B{top + 2}", "WACC ↓", italic=True, color=GREY, size=9)
        steps = (-2, -1, 0, 1, 2)
        cols = "DEFGH"
        hdr = top + 2
        for c, k in zip(cols, steps):
            s.put(f"{c}{hdr}", f"={col_ref}+({k})*{col_step}", col_fmt, bold=True, fill=MID, align="center")
        for i, k in enumerate(steps):
            rr = hdr + 1 + i
            s.put(f"C{rr}", f"={wacc_ref}+({k})*{I['w_step']}", "pct2", bold=True, fill=MID, align="center")
            for c in cols:
                s.put(f"{c}{rr}", cell(f"$C{rr}", f"{c}${hdr}"), "px", align="center")
        center = f"F{hdr + 3}"
        s.ws[center].font = s.Font(name=FONT, size=10, bold=True)
        s.ws[center].border = s.Border(*(s.Side(style="medium", color=NAVY),) * 4)
        return f"D{hdr + 1}:H{hdr + 5}", hdr

    def perp_cell(w: str, g: str) -> str:
        num = f"SUMPRODUCT({ufcf}/(1+{w})^{per})+{tfcf}*(1+{g})/({w}-{g})/(1+{w})^{tvp_perp}+{claims}"
        return f'=IFERROR(IF(OR({w}<={g},{tfcf}<=0),"NM",IF({num}<=0,"NM",({num})/{shares})),"NM")'

    def exit_cell(w: str, x: str) -> str:
        num = f"SUMPRODUCT({ufcf}/(1+{w})^{per})+{x}*{ebitda5}/(1+{w})^{dcf_mod.YEARS}+{claims}"
        return f'=IFERROR(IF({num}<=0,"NM",({num})/{shares}),"NM")'

    perp_rng, perp_hdr = grid(4, "Implied share price: perpetuity growth method", "Terminal growth rate →",
                              I["g"], I["g_step"], "pct2", perp_cell)
    exit_rng, exit_hdr = grid(14, "Implied share price: exit multiple method", "Exit EV/EBITDA multiple →",
                              I["exit"], I["x_step"], "x", exit_cell)
    # Upside vs current, perpetuity method
    s.band(24, "B", "H", "Upside / (downside) to current share price: perpetuity growth method")
    s.put("C25", "Terminal growth rate →", italic=True, color=GREY, size=9)
    s.put("B26", "WACC ↓", italic=True, color=GREY, size=9)
    for c in "DEFGH":
        s.put(f"{c}26", f"={c}{perp_hdr}", "pct2", bold=True, fill=MID, align="center")
    for i in range(5):
        rr = 27 + i
        s.put(f"C{rr}", f"=C{perp_hdr + 1 + i}", "pct2", bold=True, fill=MID, align="center")
        for c in "DEFGH":
            src = f"{c}{perp_hdr + 1 + i}"
            s.put(f"{c}{rr}", f'=IF(ISNUMBER({src}),{src}/{I["price"]}-1,"NM")', "pct", align="center")
    s.ws["F29"].font = s.Font(name=FONT, size=10, bold=True)
    s.ws["F29"].border = s.Border(*(s.Side(style="medium", color=NAVY),) * 4)
    s.put("B33", "Base case boxed. NM = WACC not above growth, or implied equity value negative.", color=GREY, size=9)
    s.link("B34", f'="Current share price: $"&TEXT({I["price"]},"0.00")', size=9)

    from openpyxl.formatting.rule import CellIsRule
    up_fill = s.Fill("solid", fgColor="E2EFDA")
    dn_fill = s.Fill("solid", fgColor="FCE4D6")
    for rng in (perp_rng, exit_rng):
        s.ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[I["price"]],
                                                        fill=up_fill))
        s.ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[I["price"]], fill=dn_fill))
    s.ws.conditional_formatting.add("D27:H31", CellIsRule(operator="greaterThanOrEqual", formula=["0"],
                                                          fill=up_fill))
    s.ws.conditional_formatting.add("D27:H31", CellIsRule(operator="lessThan", formula=["0"], fill=dn_fill))
    sens.landscape("A1:H34")

    # ======================= Cover =======================
    s = cover
    s.title(f"{t.name}", f"{t.ticker} · Discounted cash flow analysis · {unit_line}", "H")
    s.widths(B=36, C=14, D=14, E=4, F=36, G=14, H=14)
    s.put("B4", f"Share price as of {price_date}; financials through {m.ltm_end} "
                f"({t.latest_filing or 'latest filing'}).", italic=True, color=GREY, size=9)

    s.band(6, "B", "D", "Valuation output", {"C": "Perpetuity", "D": "Exit multiple"})
    cov = [("Implied share price", "px", "px"), ("Upside / (downside)", "up", "pct"),
           ("Enterprise value ($m)", "ev", "usd"), ("Equity value ($m)", "eq", "usd"),
           ("Terminal value as % of EV", "tv_share", "pct")]
    rr = 6
    for label, key, fmt in cov:
        rr += 1
        s.put(f"B{rr}", label, bold=key == "px")
        s.link(f"C{rr}", f"=DCF!C{V[key]}", fmt, bold=key == "px")
        s.link(f"D{rr}", f"=DCF!D{V[key]}", fmt, bold=key == "px")
    s.shade("C7:D7", MID)

    s.band(6, "F", "H", "Key assumptions", {"G": "", "H": ""})
    keys = [("Current share price", I["price"], "px"), ("WACC", wacc_ref, "pct2"),
            ("Cost of equity", f"WACC!$D${W['ke']}", "pct2"), ("Terminal growth", I["g"], "pct2"),
            ("Exit EV/EBITDA", I["exit"], "x"), ("Year-1 revenue growth", I["growth_1"], "pct"),
            ("EBIT margin", I["ebit_margin"], "pct"), ("Tax rate", I["tax_rate"], "pct")]
    for i, (label, ref, fmt) in enumerate(keys, 7):
        s.put(f"F{i}", label)
        s.link(f"G{i}", f"={ref}", fmt)

    # Valuation summary (football field)
    top = 17
    s.band(top, "B", "D", "Valuation summary ($ per share)", {"C": "Low", "D": "High"})
    ff = [("DCF: perpetuity (WACC ±1%, g ±0.5%)",
           f"=MIN(Sensitivity!E{perp_hdr + 1}:G{perp_hdr + 5})", f"=MAX(Sensitivity!E{perp_hdr + 1}:G{perp_hdr + 5})"),
          ("DCF: exit multiple (WACC ±1%, ±1.0x)",
           f"=MIN(Sensitivity!E{exit_hdr + 1}:G{exit_hdr + 5})", f"=MAX(Sensitivity!E{exit_hdr + 1}:G{exit_hdr + 5})")]
    for bar in rep.comps.football_field():
        if bar["label"] == "Current price" or bar["low"] is None or bar["high"] is None:
            continue
        ff.append((("Comps: " if bar["label"] != "52-week range" else "") + bar["label"], bar["low"], bar["high"]))
    rr = top
    for label, lo, hi in ff:
        rr += 1
        s.put(f"B{rr}", label)
        for col, v in (("C", lo), ("D", hi)):
            if isinstance(v, str):
                s.link(f"{col}{rr}", v, "px")
            else:
                s.inp(f"{col}{rr}", v, "px", fill=False)
        s.put(f"E{rr}", f"=D{rr}-C{rr}", "px", color="FFFFFF", size=1)  # bar length for the chart (hidden)
    ff_last = rr
    rr += 1
    s.put(f"B{rr}", "Current share price", italic=True)
    s.link(f"C{rr}", f"={I['price']}", "px", italic=True)
    s.put(f"B{rr + 1}", "Comps rows: target's LTM metric × peer 25th-75th percentile multiple, bridged to equity "
                        "and divided by shares (FD-8). DCF rows: corners of the Sensitivity grids.",
          color=GREY, size=8, wrap=False)

    from openpyxl.chart import BarChart, Reference
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
    ch.title = None
    ch.y_axis.title = "Implied share price ($)"
    ch.y_axis.numFmt = '$#,##0'
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"  # first row on top
    ch.legend = None
    cats = Reference(s.ws, min_col=2, min_row=top + 1, max_row=ff_last)
    lows = Reference(s.ws, min_col=3, min_row=top + 1, max_row=ff_last)
    spans = Reference(s.ws, min_col=5, min_row=top + 1, max_row=ff_last)
    ch.add_data(lows, titles_from_data=False)
    ch.add_data(spans, titles_from_data=False)
    ch.set_categories(cats)
    ch.series[0].graphicalProperties.noFill = True
    ch.series[0].graphicalProperties.line.noFill = True
    ch.series[1].graphicalProperties.solidFill = NAVY
    ch.series[1].graphicalProperties.line.solidFill = NAVY
    ch.gapWidth = 60
    ch.height, ch.width = 7.5, 12
    s.ws.add_chart(ch, f"F{top}")

    rr = max(rr + 3, top + 17)
    s.section(rr, "B", "H", "How to use this file")
    for line in (
            "1. Change any blue cell on yellow on the Inputs sheet (growth, margins, WACC inputs, terminal "
            "assumptions). Every other number is a formula and updates.",
            "2. DCF sheet: historicals → projections → terminal year → enterprise value → equity value → share price, "
            "by both perpetuity growth and exit multiple.",
            "3. WACC sheet: CAPM cost of equity and after-tax cost of debt at market weights. Sensitivity sheet: "
            "price across WACC and terminal assumptions.",
            "4. Sources sheet: the XBRL tag, SEC filing and formula behind every hard-coded number. Nothing here "
            "was estimated by an AI.",
            "Colour code: blue = hard-coded input or filing figure, black = formula, green = link to another sheet."):
        rr += 1
        s.put(f"B{rr}", line, size=9)
    rr += 2
    s.put(f"B{rr}", "Model check:", bold=True, size=9)
    s.link(f"C{rr}", f"=DCF!C{check_row}", size=9)
    cover.landscape(f"A1:H{rr}")

    model.book.sources_sheet()
    src = model.wb["Sources"]
    src.page_setup.orientation = "landscape"
    src.page_setup.fitToWidth, src.page_setup.fitToHeight = 1, 0
    src.sheet_properties.pageSetUpPr.fitToPage = True
    return _save(model)


def _save(model: _Model) -> bytes:
    """Arial everywhere (keeping each cell's size, colour and weight), Cover first and active."""
    for ws in model.wb.worksheets:
        for rowc in ws.iter_rows():
            for c in rowc:
                if c.value is not None and c.font.name != FONT:
                    f = copy(c.font)
                    f.name = FONT
                    c.font = f
    model.wb.active = 0
    return model.book.save()
